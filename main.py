import io
import os
import re
import tempfile
from datetime import datetime, timedelta

import requests
from flask import (Flask, render_template, request, redirect, url_for,
                   flash, send_file, abort)
from flask_httpauth import HTTPBasicAuth
from sqlalchemy import create_engine, text, inspect, func, or_

from models import db, Client, Vehicle, Product, Order, OrderItem, StockMovement

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'bydhelp-zap-change-me')

# ---------------------------------------------------------------------------
# БАЗА ДАННЫХ
# На Render: переменная DATABASE_URL (PostgreSQL, например Neon).
# Без неё (локально) используется файл crm.db.
# ---------------------------------------------------------------------------
DATABASE_URL = os.environ.get('DATABASE_URL', '').strip()
if DATABASE_URL.startswith('postgres://'):
    DATABASE_URL = DATABASE_URL.replace('postgres://', 'postgresql://', 1)
USE_POSTGRES = DATABASE_URL.startswith('postgresql')

app.config['SQLALCHEMY_DATABASE_URI'] = DATABASE_URL if USE_POSTGRES else 'sqlite:///crm.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'pool_pre_ping': True, 'pool_recycle': 280}
app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024  # Excel до 10 МБ

db.init_app(app)
auth = HTTPBasicAuth()

USERS = {
    os.environ.get('ADMIN_USER', 'admin'): os.environ.get('ADMIN_PASSWORD', 'zap12345')
}


@auth.verify_password
def verify_password(username, password):
    if username in USERS and USERS[username] == password:
        return username


# ---------------------------------------------------------------------------
# СПРАВОЧНИКИ
# ---------------------------------------------------------------------------
TZ_OFFSET = timedelta(hours=5)  # Ташкент, UTC+5

ORDER_STATUSES = [
    'Выдано со склада',
    'Заказ оплачен',
    'Согласовано с клиентом',
    'Идет подбор аналогов',
    'Запрос поступил (под заказ)',
]
REQUEST_STATUSES = ORDER_STATUSES[2:]
PAYMENT_METHODS = ['Наличные', 'Терминал', 'Перевод на карту']
WRITEOFF_REASONS = ['Брак', 'Пересорт', 'Потеря', 'Использовано в сервисе', 'Другое']

KIND_IN = 'Приход'
KIND_SALE = 'Продажа'
KIND_RETURN = 'Возврат'
KIND_WRITEOFF = 'Списание'
KIND_ADJUST = 'Корректировка'
MOVEMENT_KINDS = [KIND_IN, KIND_SALE, KIND_RETURN, KIND_WRITEOFF, KIND_ADJUST]


# ---------------------------------------------------------------------------
# ПОМОЩНИКИ
# ---------------------------------------------------------------------------
def to_float(value, default=0.0):
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(' ', '').replace(' ', '').replace('$', '').replace(',', '.')
    if not s:
        return default
    try:
        return float(s)
    except ValueError:
        return default


def to_int(value, default=0):
    return int(round(to_float(value, default)))


def local_now():
    return datetime.utcnow() + TZ_OFFSET


def local_date_bounds(date_from, date_to):
    """Локальные даты (YYYY-MM-DD) -> границы в UTC: [начало, конец)."""
    start = datetime.strptime(date_from, '%Y-%m-%d') - TZ_OFFSET
    end = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1) - TZ_OFFSET
    return start, end


def month_period():
    now = local_now()
    first = now.replace(day=1)
    return first.strftime('%Y-%m-%d'), now.strftime('%Y-%m-%d')


def read_period():
    d_from = request.args.get('from') or ''
    d_to = request.args.get('to') or ''
    default_from, default_to = month_period()
    try:
        datetime.strptime(d_from, '%Y-%m-%d')
    except ValueError:
        d_from = default_from
    try:
        datetime.strptime(d_to, '%Y-%m-%d')
    except ValueError:
        d_to = default_to
    return d_from, d_to


def log_movement(product, kind, qty, price=0.0, supplier=None, comment=None, order_id=None):
    db.session.add(StockMovement(
        product_id=product.id if product else None,
        oem=product.oem if product else None,
        product_name=product.name if product else None,
        kind=kind,
        quantity=qty,
        price=price or 0.0,
        supplier=supplier,
        comment=comment,
        order_id=order_id,
        stock_after=product.stock if product else None,
    ))


def search_products(q):
    like = f'%{q}%'
    return Product.query.filter(or_(
        Product.oem.ilike(like),
        Product.name.ilike(like),
        Product.brand.ilike(like),
        Product.analogs.ilike(like),
        Product.location.ilike(like),
    ))


def item_cost(item, purchase_by_oem):
    if item.purchase_price is not None:
        return item.purchase_price * item.quantity
    return purchase_by_oem.get(item.oem, 0.0) * item.quantity


@app.template_filter('money')
def money_filter(value):
    try:
        v = float(value or 0)
    except (TypeError, ValueError):
        v = 0.0
    s = f'{v:,.2f}'.replace(',', ' ')
    if s.endswith('.00'):
        s = s[:-3]
    return '$' + s


@app.template_filter('local')
def local_filter(dt, fmt='%d.%m.%Y %H:%M'):
    if not dt:
        return ''
    return (dt + TZ_OFFSET).strftime(fmt)


@app.context_processor
def inject_globals():
    return {
        'USE_POSTGRES': USE_POSTGRES,
        'ORDER_STATUSES': ORDER_STATUSES,
        'PAYMENT_METHODS': PAYMENT_METHODS,
        'nav': request.endpoint,
    }


# ---------------------------------------------------------------------------
# ЯНДЕКС.ДИСК (перенос старых данных / бэкап в режиме SQLite)
# ---------------------------------------------------------------------------
YANDEX_TOKEN = os.environ.get('YANDEX_TOKEN', '').strip()
BACKUP_PATH_ON_DISK = 'app:/crm_backup.db'


def download_yandex_backup(dest_path):
    if not YANDEX_TOKEN:
        return False
    headers = {'Authorization': f'OAuth {YANDEX_TOKEN}'}
    try:
        url = f"https://cloud-api.yandex.net/v1/disk/resources/download?path={BACKUP_PATH_ON_DISK}"
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200:
            r = requests.get(response.json().get('href'), timeout=60)
            if r.status_code == 200 and r.content:
                with open(dest_path, 'wb') as f:
                    f.write(r.content)
                return True
    except Exception as e:
        print(f"Ошибка загрузки бэкапа: {e}")
    return False


def backup_to_yandex_disk():
    if USE_POSTGRES or not YANDEX_TOKEN:
        return
    db_path = os.path.join(app.instance_path, 'crm.db')
    if not os.path.exists(db_path):
        return
    headers = {'Authorization': f'OAuth {YANDEX_TOKEN}'}
    try:
        url = f"https://cloud-api.yandex.net/v1/disk/resources/upload?path={BACKUP_PATH_ON_DISK}&overwrite=true"
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code == 200:
            with open(db_path, 'rb') as f:
                requests.put(response.json().get('href'), data=f, timeout=60)
    except Exception as e:
        print(f"Ошибка бэкапа: {e}")


def restore_sqlite_from_yandex():
    if USE_POSTGRES or not YANDEX_TOKEN:
        return
    os.makedirs(app.instance_path, exist_ok=True)
    download_yandex_backup(os.path.join(app.instance_path, 'crm.db'))


OLD_MODELS = [Client, Vehicle, Product, Order, OrderItem]


def import_from_sqlite_file(sqlite_path):
    """Переносит данные из старого файла SQLite в текущую базу."""
    src_engine = create_engine(f'sqlite:///{sqlite_path}')
    total = 0
    try:
        with src_engine.connect() as src:
            src_insp = inspect(src)
            src_tables = set(src_insp.get_table_names())
            for model in OLD_MODELS:
                table = model.__table__
                if table.name not in src_tables:
                    continue
                src_cols = {c['name'] for c in src_insp.get_columns(table.name)}
                cols = [c for c in table.columns if c.name in src_cols]
                rows = [dict(r._mapping) for r in src.execute(
                    text('SELECT {} FROM "{}"'.format(', '.join(f'"{c.name}"' for c in cols), table.name))
                )]
                # даты из SQLite приходят строками
                for row in rows:
                    for key, val in row.items():
                        if isinstance(val, str) and key == 'created_at':
                            try:
                                row[key] = datetime.fromisoformat(val)
                            except ValueError:
                                row[key] = datetime.utcnow()
                if rows:
                    db.session.execute(table.insert(), rows)
                    total += len(rows)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        print(f"Ошибка переноса данных: {e}")
        return 0
    finally:
        src_engine.dispose()
    reset_sequences()
    return total


def reset_sequences():
    if not USE_POSTGRES:
        return
    for model in OLD_MODELS + [StockMovement]:
        t = model.__table__.name
        db.session.execute(text(
            f"SELECT setval(pg_get_serial_sequence('\"{t}\"', 'id'), "
            f"COALESCE((SELECT MAX(id) FROM \"{t}\"), 0) + 1, false)"
        ))
    db.session.commit()


def migrate_old_data_if_needed():
    if not USE_POSTGRES or not YANDEX_TOKEN:
        return
    if Product.query.first() or Order.query.first():
        return
    tmp_path = os.path.join(tempfile.gettempdir(), 'crm_backup_import.db')
    if download_yandex_backup(tmp_path):
        print(f"Перенесено записей из старой базы: {import_from_sqlite_file(tmp_path)}")
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def ensure_columns():
    """Добавляет в существующие таблицы новые колонки (без потери данных)."""
    insp = inspect(db.engine)
    prep = db.engine.dialect.identifier_preparer
    with db.engine.begin() as conn:
        for table in db.metadata.sorted_tables:
            if not insp.has_table(table.name):
                continue
            existing = {c['name'] for c in insp.get_columns(table.name)}
            for col in table.columns:
                if col.name in existing:
                    continue
                col_type = col.type.compile(dialect=db.engine.dialect)
                conn.execute(text(
                    f'ALTER TABLE {prep.quote(table.name)} ADD COLUMN {prep.quote(col.name)} {col_type}'
                ))
                print(f'Добавлена колонка {table.name}.{col.name}')


with app.app_context():
    restore_sqlite_from_yandex()
    db.create_all()
    ensure_columns()
    migrate_old_data_if_needed()


# ---------------------------------------------------------------------------
# ГЛАВНАЯ
# ---------------------------------------------------------------------------
@app.route('/')
@auth.login_required
def index():
    d_from, d_to = month_period()
    start, end = local_date_bounds(d_from, d_to)

    products = Product.query.all()
    purchase_by_oem = {p.oem: (p.purchase_price or 0) for p in products}

    orders = Order.query.filter(Order.created_at >= start, Order.created_at < end).all()
    sold_orders = [o for o in orders if o.items]
    sales = sum(o.total_price or 0 for o in sold_orders)
    cost = sum(item_cost(i, purchase_by_oem) for o in sold_orders for i in o.items)

    low = sorted([p for p in products if p.is_low], key=lambda p: (p.stock or 0))
    open_requests = Order.query.filter(Order.status.in_(REQUEST_STATUSES)).order_by(Order.id.desc()).limit(6).all()
    movements = StockMovement.query.order_by(StockMovement.id.desc()).limit(8).all()

    stats = {
        'sales': sales,
        'profit': sales - cost,
        'orders': len(sold_orders),
        'avg': (sales / len(sold_orders)) if sold_orders else 0,
        'stock_cost': sum((p.purchase_price or 0) * max(p.stock or 0, 0) for p in products),
        'stock_sell': sum((p.sell_price or 0) * max(p.stock or 0, 0) for p in products),
        'positions': len(products),
        'units': sum(max(p.stock or 0, 0) for p in products),
        'low': len(low),
    }
    return render_template('index.html', stats=stats, low=low[:8], open_requests=open_requests,
                           movements=movements, month_label=local_now().strftime('%m.%Y'))


# ---------------------------------------------------------------------------
# СКЛАД
# ---------------------------------------------------------------------------
@app.route('/products')
@auth.login_required
def products_page():
    q = (request.args.get('q') or '').strip()
    flt = request.args.get('filter', '')
    query = search_products(q) if q else Product.query
    if flt == 'out':
        query = query.filter(Product.stock <= 0)
    products = query.order_by(Product.name).all()
    if flt == 'low':
        products = [p for p in products if p.is_low]
    counts = {
        'all': Product.query.count(),
        'out': Product.query.filter(Product.stock <= 0).count(),
        'low': sum(1 for p in Product.query.all() if p.is_low),
    }
    return render_template('products.html', products=products, query=q, flt=flt, counts=counts)


@app.route('/add_product', methods=['POST'])
@auth.login_required
def add_product():
    oem = (request.form.get('oem') or '').strip()
    name = (request.form.get('name') or '').strip()
    if not oem or not name:
        flash('Укажите OEM и название', 'error')
        return redirect(url_for('products_page'))
    if Product.query.filter(func.lower(Product.oem) == oem.lower()).first():
        flash(f'Запчасть с OEM {oem} уже есть в базе — оформите приход на неё', 'error')
        return redirect(url_for('products_page', q=oem))

    stock = to_int(request.form.get('stock'), 0)
    p = Product(
        oem=oem, name=name,
        brand=(request.form.get('brand') or '').strip(),
        purchase_price=to_float(request.form.get('purchase_price')),
        sell_price=to_float(request.form.get('sell_price')),
        stock=stock,
        location=(request.form.get('location') or '').strip(),
        min_stock=to_int(request.form.get('min_stock'), 0),
        analogs=(request.form.get('analogs') or '').strip(),
    )
    db.session.add(p)
    db.session.flush()
    if stock:
        log_movement(p, KIND_IN, stock, p.purchase_price,
                     supplier=(request.form.get('supplier') or '').strip() or None,
                     comment='Начальный остаток')
    db.session.commit()
    flash(f'Добавлено: {p.name}', 'ok')
    return redirect(url_for('product_page', id=p.id))


@app.route('/product/<int:id>')
@auth.login_required
def product_page(id):
    p = db.session.get(Product, id) or abort(404)
    history = StockMovement.query.filter_by(product_id=id).order_by(StockMovement.id.desc()).limit(200).all()
    sold_qty = db.session.query(func.coalesce(func.sum(OrderItem.quantity), 0)).filter(OrderItem.oem == p.oem).scalar()
    return render_template('product.html', p=p, history=history, sold_qty=sold_qty,
                           reasons=WRITEOFF_REASONS)


@app.route('/edit_product/<int:id>', methods=['POST'])
@auth.login_required
def edit_product(id):
    p = db.session.get(Product, id) or abort(404)
    new_oem = (request.form.get('oem') or p.oem).strip()
    if new_oem.lower() != p.oem.lower() and Product.query.filter(func.lower(Product.oem) == new_oem.lower()).first():
        flash(f'OEM {new_oem} уже занят другой запчастью', 'error')
        return redirect(url_for('product_page', id=id))
    p.oem = new_oem
    p.name = (request.form.get('name') or p.name).strip()
    p.brand = (request.form.get('brand') or '').strip()
    p.purchase_price = to_float(request.form.get('purchase_price'), p.purchase_price or 0)
    p.sell_price = to_float(request.form.get('sell_price'), p.sell_price or 0)
    p.location = (request.form.get('location') or '').strip()
    p.min_stock = to_int(request.form.get('min_stock'), p.min_stock or 0)
    p.analogs = (request.form.get('analogs') or '').strip()
    p.note = (request.form.get('note') or '').strip()

    if request.form.get('stock') not in (None, ''):
        new_stock = to_int(request.form.get('stock'), p.stock or 0)
        diff = new_stock - (p.stock or 0)
        if diff:
            p.stock = new_stock
            log_movement(p, KIND_ADJUST, diff, comment='Ручное изменение остатка')
    db.session.commit()
    flash('Сохранено', 'ok')
    return redirect(request.form.get('next') or url_for('product_page', id=id))


@app.route('/delete_product/<int:id>', methods=['POST'])
@auth.login_required
def delete_product(id):
    p = db.session.get(Product, id) or abort(404)
    name = p.name
    db.session.delete(p)
    db.session.commit()
    flash(f'Удалено: {name}', 'ok')
    return redirect(url_for('products_page'))


# ---------------------------------------------------------------------------
# ПРИХОД / СПИСАНИЕ
# ---------------------------------------------------------------------------
@app.route('/receive', methods=['GET', 'POST'])
@auth.login_required
def receive_page():
    if request.method == 'POST':
        oem = (request.form.get('oem') or '').strip()
        qty = to_int(request.form.get('quantity'), 0)
        price = request.form.get('purchase_price')
        sell = request.form.get('sell_price')
        supplier = (request.form.get('supplier') or '').strip() or None
        comment = (request.form.get('comment') or '').strip() or None

        if not oem or qty <= 0:
            flash('Укажите OEM и количество больше нуля', 'error')
            return redirect(url_for('receive_page'))

        p = Product.query.filter(func.lower(Product.oem) == oem.lower()).first()
        if not p:
            name = (request.form.get('name') or '').strip()
            if not name:
                flash(f'OEM {oem} нет в базе — укажите название, чтобы создать новую запчасть', 'error')
                return redirect(url_for('receive_page', oem=oem))
            p = Product(oem=oem, name=name, brand=(request.form.get('brand') or '').strip(),
                        stock=0, purchase_price=0, sell_price=0,
                        location=(request.form.get('location') or '').strip(), min_stock=0)
            db.session.add(p)
            db.session.flush()

        if price not in (None, ''):
            p.purchase_price = to_float(price, p.purchase_price or 0)
        if sell not in (None, ''):
            p.sell_price = to_float(sell, p.sell_price or 0)
        if request.form.get('location'):
            p.location = request.form.get('location').strip()
        p.stock = (p.stock or 0) + qty
        log_movement(p, KIND_IN, qty, p.purchase_price, supplier=supplier, comment=comment)
        db.session.commit()
        flash(f'Приход: {p.name} +{qty} шт. Теперь на складе {p.stock} шт.', 'ok')
        return redirect(url_for('receive_page'))

    products = Product.query.order_by(Product.name).all()
    recent = StockMovement.query.filter_by(kind=KIND_IN).order_by(StockMovement.id.desc()).limit(15).all()
    suppliers = [s for (s,) in db.session.query(StockMovement.supplier)
                 .filter(StockMovement.supplier.isnot(None)).distinct().limit(50)]
    return render_template('receive.html', products=products, recent=recent, suppliers=suppliers,
                           preset_oem=request.args.get('oem', ''))


@app.route('/writeoff/<int:id>', methods=['POST'])
@auth.login_required
def writeoff(id):
    p = db.session.get(Product, id) or abort(404)
    qty = to_int(request.form.get('quantity'), 0)
    if qty <= 0 or qty > (p.stock or 0):
        flash(f'Нельзя списать {qty} шт. — на складе {p.stock} шт.', 'error')
        return redirect(url_for('product_page', id=id))
    reason = request.form.get('reason') or 'Другое'
    comment = (request.form.get('comment') or '').strip()
    p.stock -= qty
    log_movement(p, KIND_WRITEOFF, -qty, p.purchase_price,
                 comment=f'{reason}' + (f': {comment}' if comment else ''))
    db.session.commit()
    flash(f'Списано {qty} шт.', 'ok')
    return redirect(url_for('product_page', id=id))


def safe_next(url):
    """Возврат только на страницы этого сайта."""
    if url and url.startswith('/') and not url.startswith('//'):
        return url
    return url_for('receive_page')


@app.route('/movement/<int:id>/edit', methods=['GET', 'POST'])
@auth.login_required
def edit_movement(id):
    """Исправление прихода: количество, цена, поставщик, комментарий, запчасть."""
    m = db.session.get(StockMovement, id) or abort(404)
    if m.kind != KIND_IN:
        flash('Изменять можно только приход. Продажи правятся через заказы.', 'error')
        return redirect(url_for('movements_page'))

    old_product = db.session.get(Product, m.product_id) if m.product_id else None
    back = safe_next(request.values.get('next'))

    if request.method == 'POST':
        new_oem = (request.form.get('oem') or '').strip()
        new_qty = to_int(request.form.get('quantity'), 0)
        new_price = to_float(request.form.get('price'), m.price or 0)
        if new_qty <= 0:
            flash('Количество должно быть больше нуля', 'error')
            return redirect(url_for('edit_movement', id=id, next=back))

        new_product = Product.query.filter(func.lower(Product.oem) == new_oem.lower()).first() if new_oem else old_product
        if not new_product:
            flash(f'Запчасти с OEM {new_oem} нет в базе', 'error')
            return redirect(url_for('edit_movement', id=id, next=back))

        # убираем старый приход с остатка и добавляем новый
        if old_product and old_product.id == new_product.id:
            diff = new_qty - m.quantity
            if (new_product.stock or 0) + diff < 0:
                flash(f'Нельзя уменьшить приход: часть товара уже продана. На складе {new_product.stock} шт.', 'error')
                return redirect(url_for('edit_movement', id=id, next=back))
            new_product.stock = (new_product.stock or 0) + diff
        else:
            if old_product:
                if (old_product.stock or 0) - m.quantity < 0:
                    flash(f'Нельзя перенести приход: {old_product.name} уже продан(а). На складе {old_product.stock} шт.', 'error')
                    return redirect(url_for('edit_movement', id=id, next=back))
                old_product.stock = (old_product.stock or 0) - m.quantity
            new_product.stock = (new_product.stock or 0) + new_qty

        m.product_id = new_product.id
        m.oem = new_product.oem
        m.product_name = new_product.name
        m.quantity = new_qty
        m.price = new_price
        m.supplier = (request.form.get('supplier') or '').strip() or None
        m.comment = (request.form.get('comment') or '').strip() or None
        m.stock_after = new_product.stock

        if request.form.get('update_price'):
            new_product.purchase_price = new_price
        sell = request.form.get('sell_price')
        if sell not in (None, ''):
            new_product.sell_price = to_float(sell, new_product.sell_price or 0)

        db.session.commit()
        flash(f'Приход исправлен: {new_product.name}, {new_qty} шт. по {money_filter(new_price)}. На складе {new_product.stock} шт.', 'ok')
        return redirect(back)

    products = Product.query.order_by(Product.name).all()
    suppliers = [s for (s,) in db.session.query(StockMovement.supplier)
                 .filter(StockMovement.supplier.isnot(None)).distinct().limit(50)]
    return render_template('movement_edit.html', m=m, p=old_product, products=products,
                           suppliers=suppliers, back=back)


@app.route('/movement/<int:id>/delete', methods=['POST'])
@auth.login_required
def delete_movement(id):
    """Отмена ошибочного прихода: товар убирается с остатка, запись удаляется."""
    m = db.session.get(StockMovement, id) or abort(404)
    back = safe_next(request.form.get('next'))
    if m.kind != KIND_IN:
        flash('Удалять можно только приход', 'error')
        return redirect(back)
    p = db.session.get(Product, m.product_id) if m.product_id else None
    if p:
        if (p.stock or 0) - m.quantity < 0:
            flash(f'Нельзя отменить приход: часть товара уже продана. На складе {p.stock} шт.', 'error')
            return redirect(back)
        p.stock = (p.stock or 0) - m.quantity
    db.session.delete(m)
    db.session.commit()
    flash('Приход отменён' + (f', на складе {p.name} теперь {p.stock} шт.' if p else ''), 'ok')
    return redirect(back)


# ---------------------------------------------------------------------------
# ПРОДАЖА / ЗАКАЗЫ
# ---------------------------------------------------------------------------
@app.route('/sale')
@auth.login_required
def sale_page():
    products = Product.query.order_by(Product.name).all()
    catalog = [{
        'id': p.id, 'oem': p.oem, 'name': p.name, 'brand': p.brand or '',
        'price': p.sell_price or 0, 'stock': p.stock or 0, 'loc': p.location or '',
        'analogs': p.analogs or '',
    } for p in products]
    return render_template('sale.html', catalog=catalog)


@app.route('/create_order', methods=['POST'])
@auth.login_required
def create_order():
    first_name = (request.form.get('first_name') or '').strip()
    last_name = (request.form.get('last_name') or '').strip()
    phone = (request.form.get('phone') or '').strip()
    vin = (request.form.get('vin') or '').strip().upper()
    comment = (request.form.get('comment') or '').strip()
    status = request.form.get('status') or ORDER_STATUSES[0]
    payment = request.form.get('payment_method') or None
    discount = max(to_float(request.form.get('discount')), 0.0)

    ids = request.form.getlist('item_id')
    qtys = request.form.getlist('item_qty')
    prices = request.form.getlist('item_price')

    # собираем позиции (одинаковые товары складываем)
    lines = {}
    for i, raw_id in enumerate(ids):
        pid = to_int(raw_id, 0)
        qty = to_int(qtys[i] if i < len(qtys) else 0, 0)
        if pid <= 0 or qty <= 0:
            continue
        price = to_float(prices[i] if i < len(prices) else None, -1)
        if pid in lines:
            lines[pid]['qty'] += qty
        else:
            lines[pid] = {'qty': qty, 'price': price}

    if not lines and not comment:
        flash('Добавьте товары в чек или опишите запрос в комментарии', 'error')
        return redirect(url_for('sale_page'))

    products = {p.id: p for p in Product.query.filter(Product.id.in_(lines.keys())).all()} if lines else {}
    problems = []
    for pid, line in lines.items():
        p = products.get(pid)
        if not p:
            problems.append(f'товар #{pid} не найден')
        elif line['qty'] > (p.stock or 0):
            problems.append(f'{p.name} ({p.oem}): нужно {line["qty"]}, на складе {p.stock}')
    if problems:
        flash('Не хватает на складе: ' + '; '.join(problems), 'error')
        return redirect(url_for('sale_page'))

    if not lines and status in ORDER_STATUSES[:2]:
        status = REQUEST_STATUSES[-1]

    order = Order(
        first_name=first_name or 'Розничный покупатель',
        last_name=last_name,
        phone=phone or '—',
        vin=vin or None,
        comment=comment or None,
        status=status,
        payment_method=payment if lines else None,
        discount=discount if lines else 0.0,
        total_price=0.0,
    )
    db.session.add(order)
    db.session.flush()

    total = 0.0
    for pid, line in lines.items():
        p = products[pid]
        price = line['price'] if line['price'] >= 0 else (p.sell_price or 0)
        total += price * line['qty']
        db.session.add(OrderItem(
            order_id=order.id, product_id=p.id, product_name=p.name, oem=p.oem,
            price=price, quantity=line['qty'], purchase_price=p.purchase_price or 0,
        ))
        p.stock -= line['qty']
        log_movement(p, KIND_SALE, -line['qty'], price, order_id=order.id,
                     comment=f'Заказ №{order.id}')

    order.total_price = max(total - order.discount, 0.0)
    db.session.commit()

    if lines:
        flash(f'Продажа №{order.id} оформлена на {money_filter(order.total_price)}', 'ok')
        return redirect(url_for('orders_page', open=order.id))
    flash(f'Запрос №{order.id} сохранён', 'ok')
    return redirect(url_for('orders_page'))


@app.route('/orders')
@auth.login_required
def orders_page():
    q = (request.args.get('q') or '').strip()
    st = request.args.get('status', '')
    query = Order.query
    if q:
        like = f'%{q}%'
        query = query.filter(or_(Order.first_name.ilike(like), Order.last_name.ilike(like),
                                 Order.phone.ilike(like), Order.vin.ilike(like),
                                 Order.comment.ilike(like)))
    if st == 'requests':
        query = query.filter(Order.status.in_(REQUEST_STATUSES))
    elif st == 'sales':
        query = query.filter(Order.status.in_(ORDER_STATUSES[:2]))
    orders = query.order_by(Order.id.desc()).limit(300).all()
    return render_template('orders.html', orders=orders, query=q, st=st,
                           open_id=request.args.get('open', type=int))


@app.route('/order/<int:id>/update', methods=['POST'])
@auth.login_required
def update_order(id):
    o = db.session.get(Order, id) or abort(404)
    if request.form.get('status'):
        o.status = request.form.get('status')
    if 'payment_method' in request.form:
        o.payment_method = request.form.get('payment_method') or None
    if 'comment' in request.form:
        o.comment = (request.form.get('comment') or '').strip() or None
    db.session.commit()
    flash(f'Заказ №{o.id} обновлён', 'ok')
    return redirect(url_for('orders_page', open=o.id, q=request.form.get('q', ''), status=request.form.get('st', '')))


@app.route('/delete_order/<int:id>', methods=['POST'])
@auth.login_required
def delete_order(id):
    o = db.session.get(Order, id) or abort(404)
    returned = 0
    for item in o.items:
        p = db.session.get(Product, item.product_id) if item.product_id else None
        if not p:
            p = Product.query.filter_by(oem=item.oem).first()
        if p:
            p.stock = (p.stock or 0) + item.quantity
            returned += item.quantity
            log_movement(p, KIND_RETURN, item.quantity, item.price, order_id=o.id,
                         comment=f'Удалён заказ №{o.id}')
    db.session.delete(o)
    db.session.commit()
    flash(f'Заказ №{id} удалён' + (f', на склад возвращено {returned} шт.' if returned else ''), 'ok')
    return redirect(url_for('orders_page'))


@app.route('/order/<int:id>/receipt')
@auth.login_required
def receipt(id):
    o = db.session.get(Order, id) or abort(404)
    subtotal = sum(i.price * i.quantity for i in o.items)
    return render_template('receipt.html', o=o, subtotal=subtotal)


# ---------------------------------------------------------------------------
# ЖУРНАЛ ДВИЖЕНИЯ
# ---------------------------------------------------------------------------
@app.route('/movements')
@auth.login_required
def movements_page():
    d_from, d_to = read_period()
    start, end = local_date_bounds(d_from, d_to)
    kind = request.args.get('kind', '')
    q = (request.args.get('q') or '').strip()
    query = StockMovement.query.filter(StockMovement.created_at >= start, StockMovement.created_at < end)
    if kind in MOVEMENT_KINDS:
        query = query.filter(StockMovement.kind == kind)
    if q:
        like = f'%{q}%'
        query = query.filter(or_(StockMovement.oem.ilike(like), StockMovement.product_name.ilike(like),
                                 StockMovement.supplier.ilike(like), StockMovement.comment.ilike(like)))
    rows = query.order_by(StockMovement.id.desc()).limit(1000).all()
    totals = {}
    for r in rows:
        t = totals.setdefault(r.kind, {'qty': 0, 'sum': 0.0})
        t['qty'] += abs(r.quantity)
        t['sum'] += abs(r.quantity) * (r.price or 0)
    return render_template('movements.html', rows=rows, d_from=d_from, d_to=d_to, kind=kind,
                           q=q, kinds=MOVEMENT_KINDS, totals=totals)


# ---------------------------------------------------------------------------
# ОТЧЁТЫ
# ---------------------------------------------------------------------------
@app.route('/reports')
@auth.login_required
def reports_page():
    d_from, d_to = read_period()
    start, end = local_date_bounds(d_from, d_to)

    products = Product.query.all()
    purchase_by_oem = {p.oem: (p.purchase_price or 0) for p in products}
    orders = Order.query.filter(Order.created_at >= start, Order.created_at < end).all()
    sold = [o for o in orders if o.items]

    sales = sum(o.total_price or 0 for o in sold)
    cost = sum(item_cost(i, purchase_by_oem) for o in sold for i in o.items)
    discounts = sum(o.discount or 0 for o in sold)

    by_payment = {}
    for o in sold:
        key = o.payment_method or 'Не указано'
        by_payment.setdefault(key, {'count': 0, 'sum': 0.0})
        by_payment[key]['count'] += 1
        by_payment[key]['sum'] += o.total_price or 0

    top = {}
    for o in sold:
        for i in o.items:
            t = top.setdefault(i.oem, {'oem': i.oem, 'name': i.product_name, 'qty': 0, 'sum': 0.0, 'profit': 0.0})
            t['qty'] += i.quantity
            t['sum'] += i.price * i.quantity
            t['profit'] += i.price * i.quantity - item_cost(i, purchase_by_oem)
    top_list = sorted(top.values(), key=lambda x: x['sum'], reverse=True)[:15]

    # по дням
    daily = {}
    for o in sold:
        day = (o.created_at + TZ_OFFSET).strftime('%d.%m')
        daily[day] = daily.get(day, 0) + (o.total_price or 0)
    day_cursor = datetime.strptime(d_from, '%Y-%m-%d')
    day_end = datetime.strptime(d_to, '%Y-%m-%d')
    days = []
    while day_cursor <= day_end and len(days) < 93:
        key = day_cursor.strftime('%d.%m')
        days.append({'label': key, 'sum': daily.get(key, 0)})
        day_cursor += timedelta(days=1)
    max_day = max([d['sum'] for d in days] + [1])

    # неликвид: есть остаток, но не продавалось 90 дней
    since = datetime.utcnow() - timedelta(days=90)
    recent_oems = {oem for (oem,) in db.session.query(OrderItem.oem).join(Order)
                   .filter(Order.created_at >= since).distinct()}
    dead = sorted([p for p in products if (p.stock or 0) > 0 and p.oem not in recent_oems],
                  key=lambda p: (p.purchase_price or 0) * (p.stock or 0), reverse=True)[:20]
    dead_value = sum((p.purchase_price or 0) * (p.stock or 0) for p in dead)

    stats = {
        'sales': sales, 'cost': cost, 'profit': sales - cost, 'discounts': discounts,
        'orders': len(sold), 'avg': sales / len(sold) if sold else 0,
        'margin': ((sales - cost) / sales * 100) if sales else 0,
        'requests': sum(1 for o in orders if not o.items),
    }
    return render_template('reports.html', stats=stats, d_from=d_from, d_to=d_to,
                           by_payment=by_payment, top=top_list, days=days, max_day=max_day,
                           dead=dead, dead_value=dead_value)


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------
EXPORT_HEADERS = ['OEM', 'Название', 'Бренд', 'Закуп', 'Продажа', 'Остаток', 'Полка', 'Мин. остаток', 'Аналоги']

HEADER_ALIASES = {
    'oem': ['oem', 'артикул', 'oemартикул', 'номер', 'каталожныйномер', 'код'],
    'name': ['название', 'наименование', 'name', 'деталь'],
    'brand': ['бренд', 'производитель', 'brand', 'марка'],
    'purchase_price': ['закуп', 'закупка', 'закупочнаяцена', 'ценазакупа', 'себестоимость', 'вход'],
    'sell_price': ['продажа', 'ценапродажи', 'цена', 'розница', 'розничнаяцена'],
    'stock': ['остаток', 'количество', 'колво', 'qty', 'остатокшт', 'наличие'],
    'location': ['полка', 'место', 'ячейка', 'местохранения'],
    'min_stock': ['мин', 'минимум', 'миностаток', 'минимальныйостаток'],
    'analogs': ['аналоги', 'кроссы', 'аналог'],
}


def normalize_header(h):
    return re.sub(r'[^a-zа-яё]', '', str(h or '').lower())


@app.route('/export/products')
@auth.login_required
def export_products():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Склад'
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F2937')
    for p in Product.query.order_by(Product.name).all():
        ws.append([p.oem, p.name, p.brand or '', p.purchase_price or 0, p.sell_price or 0,
                   p.stock or 0, p.location or '', p.min_stock or 0, p.analogs or ''])
    widths = [18, 40, 16, 10, 10, 10, 12, 12, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"sklad_{local_now().strftime('%Y-%m-%d')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/import/products', methods=['POST'])
@auth.login_required
def import_products():
    from openpyxl import load_workbook
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm')):
        flash('Выберите файл Excel (.xlsx)', 'error')
        return redirect(url_for('products_page'))
    mode = request.form.get('mode', 'set')  # set = остаток из файла, add = прибавить как приход
    try:
        wb = load_workbook(f, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        flash('Не удалось прочитать файл. Сохраните его как .xlsx и попробуйте снова', 'error')
        return redirect(url_for('products_page'))
    if not rows:
        flash('Файл пустой', 'error')
        return redirect(url_for('products_page'))

    colmap = {}
    for idx, h in enumerate(rows[0]):
        norm = normalize_header(h)
        for field, aliases in HEADER_ALIASES.items():
            if field not in colmap and norm in aliases:
                colmap[field] = idx
    if 'oem' not in colmap:
        flash('В первой строке нужна колонка «OEM» (или «Артикул»)', 'error')
        return redirect(url_for('products_page'))

    def cell(row, field):
        i = colmap.get(field)
        return row[i] if i is not None and i < len(row) else None

    existing = {p.oem.lower(): p for p in Product.query.all()}
    created = updated = skipped = 0
    for row in rows[1:]:
        oem = str(cell(row, 'oem') or '').strip()
        if oem.endswith('.0'):
            oem = oem[:-2]
        if not oem:
            continue
        p = existing.get(oem.lower())
        name = str(cell(row, 'name') or '').strip()
        if not p:
            if not name:
                skipped += 1
                continue
            p = Product(oem=oem, name=name, stock=0, purchase_price=0, sell_price=0, min_stock=0)
            db.session.add(p)
            db.session.flush()
            existing[oem.lower()] = p
            created += 1
        else:
            updated += 1
        if name:
            p.name = name
        for field in ('brand', 'location', 'analogs'):
            v = cell(row, field)
            if v not in (None, ''):
                setattr(p, field, str(v).strip())
        for field in ('purchase_price', 'sell_price'):
            v = cell(row, field)
            if v not in (None, ''):
                setattr(p, field, to_float(v, getattr(p, field) or 0))
        v = cell(row, 'min_stock')
        if v not in (None, ''):
            p.min_stock = to_int(v, p.min_stock or 0)
        v = cell(row, 'stock')
        if v not in (None, ''):
            qty = to_int(v, 0)
            if mode == 'add':
                if qty > 0:
                    p.stock = (p.stock or 0) + qty
                    log_movement(p, KIND_IN, qty, p.purchase_price, comment=f'Импорт из Excel: {f.filename}')
            else:
                diff = qty - (p.stock or 0)
                if diff:
                    p.stock = qty
                    log_movement(p, KIND_ADJUST, diff, p.purchase_price, comment=f'Импорт из Excel: {f.filename}')
    db.session.commit()
    msg = f'Импорт: новых {created}, обновлено {updated}'
    if skipped:
        msg += f', пропущено без названия {skipped}'
    flash(msg, 'ok')
    return redirect(url_for('products_page'))


@app.route('/manual_backup', methods=['POST'])
@auth.login_required
def manual_backup():
    backup_to_yandex_disk()
    flash('Резервная копия отправлена на Яндекс.Диск', 'ok')
    return redirect(url_for('index'))


@app.errorhandler(404)
def not_found(e):
    return render_template('message.html', title='Не найдено',
                           text='Такой записи нет — возможно, её уже удалили.'), 404


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
